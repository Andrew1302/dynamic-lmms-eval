"""One run_info sheet builder, parameterized by family.

Replaces add_run_info_{thinking,thinkadj,ablation,sweep}.py, which were four
copies of this layout differing in a title, a job template, some prose and -- as
it turned out -- their column widths (B was 74, 66 and 60). They also carried
three different job-dir resolution rules between them, one of which silently
ignored the env var it documented.

Layout, unchanged from those scripts:

    title
    <config rows: key, value>
    graph-size range per task x cell
    total prompts per model [x arm] x task
    grand total

The only structural variation is whether the family has *arms* (a 6-column
prompt table) or not (5 columns), which is exactly what distinguished
add_run_info_thinking from add_run_info_ablation.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Sequence

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

import _families
from _logs import find_sample_jsonls, iter_rows

TITLE_FONT = Font(bold=True, size=13, color="1F2A4A")
SECTION_FILL = PatternFill(fill_type="solid", fgColor="305496")
SECTION_FONT = Font(bold=True, color="FFFFFF")
KEY_FONT = Font(bold=True)
LEFT = Alignment(horizontal="left", vertical="center")
CENTER = Alignment(horizontal="center", vertical="center")

_TS_RE = re.compile(r"(\d{8}_\d{6})")


def _latest_paths(root: Path) -> list[Path]:
    """The newest timestamp's sample jsonls under a job dir."""
    paths = find_sample_jsonls(root)
    stamps = sorted({m.group(1) for p in paths if (m := _TS_RE.match(p.name))})
    if not stamps:
        return []
    return [p for p in paths if p.name.startswith(stamps[-1])]


def _blank_agg() -> list[int]:
    return [10**9, -1, 10**9, -1, 0]


def _accumulate(agg: list[int], row) -> None:
    agg[0] = min(agg[0], row.n_vertices)
    agg[1] = max(agg[1], row.n_vertices)
    agg[2] = min(agg[2], row.n_edges)
    agg[3] = max(agg[3], row.n_edges)
    agg[4] += 1


def size_ranges(spec: dict, results_root: Path, family: _families.Family | None = None) -> dict[tuple[str, str], list[int]]:
    """(task, cell) -> [nv_min, nv_max, ne_min, ne_max, n_graphs].

    Two sourcing modes, because they are not interchangeable:
      * "job"              one job per cell; the cell is part of the job name.
      * "constraint_value" one job for the whole sweep; the cell is a column in
                           the rows, which requires dataset_root so _logs can
                           back-fill it (it is only populated when
                           constraint_value != -1).
    """
    out: dict[tuple[str, str], list[int]] = {}
    tasks = set(spec["tasks"])

    if spec["cell_source"] == "constraint_value":
        for tmpl in spec["range_job_tmpls"]:
            if not tmpl:
                continue
            root = results_root / tmpl
            agg: dict[tuple[str, str], list[int]] = {}
            for row in iter_rows(_latest_paths(root), job_label=tmpl, dataset_root=root):
                if row.base_task not in tasks or row.variant != "direct":
                    continue
                agg.setdefault((row.base_task, str(row.constraint_value)), _blank_agg())
                _accumulate(agg[(row.base_task, str(row.constraint_value))], row)
            out.update({k: v for k, v in agg.items() if v[4]})
        return out

    for cell in spec["cells"]:
        for task in spec["tasks"]:
            # Templates are tried in order: the merged base job first, then the
            # legacy separate *_coloring_* job for pre-merge fetches.
            for tmpl in spec["range_job_tmpls"]:
                job = tmpl.format(cell=cell)
                agg = _blank_agg()
                for row in iter_rows(_latest_paths(results_root / job), job_label=job):
                    if row.base_task == task and row.variant == "direct":
                        _accumulate(agg, row)
                if agg[4]:
                    out[(task, cell)] = agg
                    break
    return out


def build_run_info(
    wb: openpyxl.Workbook,
    spec: dict,
    *,
    results_root: Path,
    extra_rows: Sequence[tuple[str, str]] = (),
    sizes: dict | None = None,
) -> None:
    """Replace or create the ``run_info`` sheet as the workbook's first tab."""
    if sizes is None:
        sizes = size_ranges(spec, results_root)
    arms = spec["arms"]

    if "run_info" in wb.sheetnames:
        del wb["run_info"]
    ws = wb.create_sheet("run_info", 0)

    r = 1
    ws.cell(r, 1, spec["title"]).font = TITLE_FONT
    r += 2
    for key, value in list(spec["config_rows"]) + list(extra_rows):
        ws.cell(r, 1, key).font = KEY_FONT
        ws.cell(r, 1).alignment = LEFT
        ws.cell(r, 2, value).alignment = LEFT
        r += 1
    r += 1

    # --- size range table ---
    ws.cell(r, 1, f"Graph-size range per task × {spec['cell_noun']} (dataset-exact)").font = KEY_FONT
    r += 1
    for j, h in enumerate(["task", spec["cell_header"], "nodes_min", "nodes_max", "edges_min", "edges_max", "n_graphs"], 1):
        c = ws.cell(r, j, h)
        c.fill, c.font, c.alignment = SECTION_FILL, SECTION_FONT, CENTER
    r += 1
    for task in spec["tasks"]:
        for cell in spec["cells"]:
            rng = sizes.get((task, cell))
            if rng is None:
                continue
            value = int(cell) if spec["cell_source"] == "constraint_value" else cell
            for j, val in enumerate([task, value, *rng], 1):
                ws.cell(r, j, val).alignment = LEFT if j == 1 else CENTER
            r += 1
    r += 1

    # --- prompt counts ---
    per = spec["samples_per_cell"] * len(spec["cells"])
    dims = "model × arm × task" if arms else "model × task"
    ws.cell(r, 1, f"Total prompts per {dims} (direct + disguise, summed over {spec['cells_plural']})").font = KEY_FONT
    r += 1
    headers = ["model", "arm", "task", "direct", "disguise", "total"] if arms else ["model", "task", "direct", "disguise", "total"]
    for j, h in enumerate(headers, 1):
        c = ws.cell(r, j, h)
        c.fill, c.font, c.alignment = SECTION_FILL, SECTION_FONT, CENTER
    r += 1
    label_cols = 3 if arms else 2
    grand = 0
    for model in spec["models"]:
        for arm in arms or (None,):
            for task in spec["tasks"]:
                row = [model, arm, task, per, per, 2 * per] if arms else [model, task, per, per, 2 * per]
                for j, val in enumerate(row, 1):
                    ws.cell(r, j, val).alignment = LEFT if j <= label_cols else CENTER
                grand += 2 * per
                r += 1
    r += 1
    scope = f"all models × arms × tasks × {spec['cells_plural']}" if arms else f"all models × tasks × {spec['cells_plural']}"
    ws.cell(r, 1, f"Grand total prompts ({scope})").font = KEY_FONT
    total_col = 6 if arms else 5
    ws.cell(r, total_col, grand).font = KEY_FONT
    ws.cell(r, total_col).alignment = CENTER

    for col, width in zip("ABCDEFG", spec["col_widths"]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
