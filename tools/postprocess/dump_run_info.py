#!/usr/bin/env python
"""Dump a workbook sheet to stable text, for before/after regression diffs.

openpyxl rewrites the whole zip on save (member timestamps, docProps, style
table ordering), so workbooks are never byte-comparable. The meaningful
contract is cell-level: values, bold, fill, alignment, number format, column
widths and sheet order.

    python tools/postprocess/dump_run_info.py <xlsx> [--sheet run_info]
"""

from __future__ import annotations

import argparse
from pathlib import Path

import openpyxl


def dump(path: Path, sheet: str = "run_info") -> list[str]:
    wb = openpyxl.load_workbook(path)
    out = [f"#sheets\t{','.join(wb.sheetnames)}"]
    if sheet not in wb.sheetnames:
        out.append(f"#missing\t{sheet}")
        return out
    ws = wb[sheet]
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            fill = cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type else ""
            out.append("\t".join([cell.coordinate, repr(cell.value), str(bool(cell.font and cell.font.bold)), str(fill), str(cell.alignment.horizontal or ""), str(cell.number_format)]))
    for col, dim in sorted(ws.column_dimensions.items()):
        if dim.width:
            out.append(f"#width\t{col}\t{dim.width}")
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx", type=Path)
    ap.add_argument("--sheet", default="run_info")
    args = ap.parse_args()
    print("\n".join(dump(args.xlsx, args.sheet)))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
