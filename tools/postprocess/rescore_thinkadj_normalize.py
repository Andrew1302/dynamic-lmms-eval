"""Rescore the thinkadj ablation with the fixed answer normalizer (no rerun).

The original ``_normalize`` in the dynamic_graph_benchmark task took the FIRST
whitespace token for yes/no tasks and the FIRST integer for coloring — which
scored InternVL's "A: Yes" prompt-echo as "a" (wrong) and mispriced verbose
worked solutions. The fixed normalizer (2026-07-14, see task utils.py) strips
answer-cue prefixes, falls back to a unique yes/no word in short answers, and
prefers the model's last explicit final-answer statement for integers.

This script re-reads the merged sample jsonls of the thinkadj job families
(base n=100 ``graph_bench_thinkadj_*`` + incremental ``graph_bench_thinkadj500inc_*``,
pooled — disjoint doc_id ranges 0-99 / 100-499) and reports, per
model x arm x task x variant, the logged accuracy vs the rescored accuracy.
Ground truth normalization is unaffected (targets are bare "Yes"/"14").

Coloring is sourced only from the dedicated ``_coloring_`` jobs (special-chi).

Output: remote_results/_batch_reports/thinkadj_rescored_<ts>.xlsx (+ _latest
copy) with tabs:
  run_info  — config, caveats (fp8-vision-suspect cells, temp-0.6 think decode,
              non-uniform nothink budgets), dataset-exact node/edge ranges
  headline  — model x arm x task x variant: n, logged acc, rescored acc, delta
  difficulty — same but split per difficulty

Usage:
    python tools/postprocess/rescore_thinkadj_normalize.py
"""

from __future__ import annotations

import importlib.util
import os
import re
import sys
import types
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _logs import find_sample_jsonls, iter_rows

REPO = Path(__file__).resolve().parents[2]
JOBS = REPO / "remote_results" / "08_abl_thinkadj" / "_jobs"
OUT_DIR = REPO / "remote_results" / "_batch_reports"

JOB_RE = re.compile(
    r"^graph_bench_thinkadj(?:500inc)?_(?P<arm>think|nothink)_"
    r"(?P<coloring>coloring_)?(?P<diff>easy|medium|hard)_"
    r"(?P<model>qwen35_4b|gemma4_e2b|internvl35_4b)$"
)
MODEL_PRETTY = {
    "qwen35_4b": "Qwen3.5-4B",
    "gemma4_e2b": "Gemma-4-E2B",
    "internvl35_4b": "InternVL3.5-4B",
}
TASK_ORDER = ["directed_connectivity", "shortest_path", "coloring"]

# Cells whose numbers rest on the image: connectivity/disguise is the only
# thinkadj cell that *requires* vision. Qwen3.5 + InternVL3.5 ran vllm online
# fp8 in this campaign, which corrupts their visual percept (see project
# memory fp8-vision-blindness) — flag their image-dependent cells.
FP8_SUSPECT_MODELS = {"qwen35_4b", "internvl35_4b"}

CAVEATS = [
    "fp8 VISION SUSPECT: Qwen3.5-4B and InternVL3.5-4B ran vllm online fp8 quantization in this "
    "campaign; forensic evidence (2026-07-14) shows their visual percept is corrupted (Qwen describes "
    "every render as 'stacked boxes'; InternVL reports the image as not visible; both at chance on "
    "the one vision-required cell, connectivity/disguise, where bf16 Gemma scores 0.97). Treat their "
    "connectivity/disguise cells — and ANY image-dependent conclusion — as blind-run numbers. "
    "Text-solvable cells (coloring, direct variants, shortest_path incl. disguise via the unique "
    "source/sink structure) are unaffected.",
    "DECODE: the think arm sampled at temperature=0.6, do_sample=true (official reasoning-mode "
    "guidance), NOT greedy. The nothink arm is greedy (temp 0). Paired-accuracy deficits in coloring "
    "match independent sampling noise (paired ~= orig x disg product).",
    "NOTHINK BUDGETS NON-UNIFORM: Qwen nothink got max_new_tokens=1024 plus a terse-answer "
    "directive; Gemma/InternVL kept the 64-token task default. Gemma's nothink shortest_path "
    "answers are mid-derivation truncations, not one-pass attempts.",
    "RESCORING: 'rescored acc' re-normalizes the logged filtered_resps with the fixed normalizer "
    "(task utils.py, 2026-07-14): strips 'A:'/'Answer:' prefixes; unique yes/no word fallback for "
    "short answers; last explicit final-answer statement for integers. Benchmarked deltas: "
    "connectivity +66/-0, coloring +160/-20, shortest_path +2/-2 (fixed/regressed rows). "
    "Truncation-empty responses still score 0 (genuinely unanswered).",
    "COVERAGE: think arm pooled n=500/cell for Qwen/Gemma (base 100 + inc 400); InternVL n=100; "
    "several nothink inc cells partial (see n column).",
]


def _load_normalize():
    """Import _normalize from the task utils without dragging in lmms_eval deps."""
    sys.modules.setdefault(
        "loguru",
        types.SimpleNamespace(logger=types.SimpleNamespace(info=lambda *a, **k: None)),
    )
    utils_path = REPO / "lmms_eval" / "tasks" / "dynamic_graph_benchmark" / "utils.py"
    spec = importlib.util.spec_from_file_location("dgb_task_utils", utils_path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod._normalize


def main() -> int:
    normalize = _load_normalize()

    # (model, arm, task, variant, diff) -> [n, old_correct, new_correct]
    cells: dict[tuple, list[int]] = defaultdict(lambda: [0, 0, 0])
    # (task, diff) -> [min_v, max_v, min_e, max_e]
    ranges: dict[tuple, list[int]] = {}
    n_jobs = 0

    for job_dir in sorted(JOBS.iterdir()):
        m = JOB_RE.match(job_dir.name)
        if not m or not job_dir.is_dir():
            continue
        n_jobs += 1
        arm, diff, model = m.group("arm"), m.group("diff"), m.group("model")
        is_coloring_job = bool(m.group("coloring"))
        jsonls = find_sample_jsonls(job_dir)
        for row in iter_rows(jsonls, job_label=job_dir.name):
                # coloring only from the dedicated special-chi jobs
                if (row.base_task == "coloring") != is_coloring_job:
                    continue
                old = row.correct
                new = int(normalize(row.response, row.base_task) == normalize(row.target, row.base_task))
                c = cells[(model, arm, row.base_task, row.variant, diff)]
                c[0] += 1
                c[1] += old
                c[2] += new
                r = ranges.setdefault((row.base_task, diff), [10**9, 0, 10**9, 0])
                if row.n_vertices:
                    r[0] = min(r[0], row.n_vertices)
                    r[1] = max(r[1], row.n_vertices)
                if row.n_edges:
                    r[2] = min(r[2], row.n_edges)
                    r[3] = max(r[3], row.n_edges)

    if not cells:
        print(f"no thinkadj sample logs found under {JOBS}", file=sys.stderr)
        return 1

    wb = openpyxl.Workbook()
    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    suspect_fill = PatternFill("solid", fgColor="FFF2CC")

    # --- run_info ---------------------------------------------------------
    ws = wb.active
    ws.title = "run_info"
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 120
    rows = [
        ("report", "thinkadj rescored with fixed answer normalizer"),
        ("generated", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("jobs read", str(n_jobs)),
        ("job families", "graph_bench_thinkadj_* (n=100) + graph_bench_thinkadj500inc_* (n=400), pooled"),
        ("models", ", ".join(MODEL_PRETTY.values())),
        ("prompt", "INCLUDE_ADJ_MATRIX=1 (text adjacency list) in BOTH arms"),
        ("think decode", "temperature=0.6, do_sample=true; budgets: Qwen 12288 native + 1024, Gemma/InternVL max_new_tokens=12288"),
        ("nothink decode", "greedy; Qwen max_new_tokens=1024 + terse directive, others 64 (task default)"),
    ]
    for k, v in rows:
        ws.append([k, v])
        ws.cell(ws.max_row, 1).font = bold
    ws.append([])
    ws.append(["caveats", ""])
    ws.cell(ws.max_row, 1).font = bold
    for cav in CAVEATS:
        ws.append(["", cav])
        ws.cell(ws.max_row, 2).alignment = wrap
    ws.append([])
    ws.append(["graph size ranges (dataset-exact, from row n_vertices/n_edges)", ""])
    ws.cell(ws.max_row, 1).font = bold
    ws.append(["task / difficulty", "n_vertices min-max | n_edges min-max"])
    ws.cell(ws.max_row, 1).font = bold
    ws.cell(ws.max_row, 2).font = bold
    for (task, diff) in sorted(ranges):
        r = ranges[(task, diff)]
        ws.append([f"{task} / {diff}", f"V {r[0]}-{r[1]} | E {r[2]}-{r[3]}"])

    # --- headline (difficulties pooled) -----------------------------------
    ws = wb.create_sheet("headline")
    hdr = ["model", "arm", "task", "variant", "n", "logged acc", "rescored acc", "delta", "flag"]
    ws.append(hdr)
    for c in ws[1]:
        c.font = bold
    pooled: dict[tuple, list[int]] = defaultdict(lambda: [0, 0, 0])
    for (model, arm, task, variant, _diff), (n, old, new) in cells.items():
        p = pooled[(model, arm, task, variant)]
        p[0] += n
        p[1] += old
        p[2] += new
    for model in MODEL_PRETTY:
        for arm in ("think", "nothink"):
            for task in TASK_ORDER:
                for variant in ("direct", "disguise"):
                    key = (model, arm, task, variant)
                    if key not in pooled:
                        continue
                    n, old, new = pooled[key]
                    suspect = (
                        model in FP8_SUSPECT_MODELS
                        and task == "directed_connectivity"
                        and variant == "disguise"
                    )
                    ws.append([
                        MODEL_PRETTY[model], arm, task, variant, n,
                        round(old / n, 4), round(new / n, 4), round((new - old) / n, 4),
                        "fp8-vision-suspect" if suspect else "",
                    ])
                    if suspect:
                        for c in ws[ws.max_row]:
                            c.fill = suspect_fill
    for col, w in zip("ABCDEFGHI", (16, 9, 22, 10, 7, 11, 13, 8, 20)):
        ws.column_dimensions[col].width = w

    # --- per-difficulty detail --------------------------------------------
    ws = wb.create_sheet("difficulty")
    ws.append(["model", "arm", "task", "variant", "difficulty", "n", "logged acc", "rescored acc", "delta"])
    for c in ws[1]:
        c.font = bold
    for (model, arm, task, variant, diff) in sorted(
        cells,
        key=lambda k: (list(MODEL_PRETTY).index(k[0]), k[1], TASK_ORDER.index(k[2]), k[3],
                       ["easy", "medium", "hard"].index(k[4])),
    ):
        n, old, new = cells[(model, arm, task, variant, diff)]
        ws.append([MODEL_PRETTY[model], arm, task, variant, diff, n,
                   round(old / n, 4), round(new / n, 4), round((new - old) / n, 4)])
    for col, w in zip("ABCDEFGHI", (16, 9, 22, 10, 10, 7, 11, 13, 8)):
        ws.column_dimensions[col].width = w

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out = OUT_DIR / f"thinkadj_rescored_{ts}.xlsx"
    wb.save(out)
    latest = OUT_DIR / "thinkadj_rescored_latest.xlsx"
    wb.save(latest)
    print(f"wrote {out}")
    print(f"wrote {latest}")

    # console summary of the biggest movers
    print("\nbiggest deltas (headline):")
    movers = sorted(pooled.items(), key=lambda kv: abs(kv[1][2] - kv[1][1]) / kv[1][0], reverse=True)[:8]
    for (model, arm, task, variant), (n, old, new) in movers:
        print(f"  {MODEL_PRETTY[model]:16s} {arm:7s} {task:22s} {variant:8s} "
              f"{old/n:.3f} -> {new/n:.3f}  (n={n})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
