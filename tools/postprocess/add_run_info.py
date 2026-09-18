#!/usr/bin/env python
"""Add or replace the run_info sheet on a report workbook.

Replaces add_run_info_{thinking,thinkadj,ablation}.py, which were three copies
of one layout with three different hand-rolled argv checks and three different
job-dir resolution rules.

    python tools/postprocess/add_run_info.py --family think --xlsx report.xlsx

``--results-root`` points at the tree holding the sample jsonls: the campaign
tree for a normal report, or a scratch assembled root for the full-report build
(which is what the wrapper used to monkeypatch a module global to achieve).

add_run_info_sweep.py is deliberately NOT replaced: the sweep family sources its
coloring rows and its conn/shortest_path rows from *different* jobs, which is a
different sourcing shape rather than different parameters. Folding it in would
mean a field with exactly one user.
"""

from __future__ import annotations

import argparse
import os
import sys
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import openpyxl

import _runinfo
from _runinfo_specs import RUN_INFO_SPECS

REPO = Path(__file__).resolve().parents[2]
DEFAULT_RESULTS_ROOT = REPO / "remote_results"


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--family", required=True, choices=sorted(RUN_INFO_SPECS))
    ap.add_argument("--xlsx", required=True, type=Path, help="workbook to modify in place")
    ap.add_argument("--results-root", type=Path, default=DEFAULT_RESULTS_ROOT, help=f"tree holding the sample jsonls (default: {DEFAULT_RESULTS_ROOT})")
    args = ap.parse_args(argv)

    if not args.xlsx.is_file():
        ap.error(f"no such workbook: {args.xlsx}")

    wb = openpyxl.load_workbook(args.xlsx)
    _runinfo.build_run_info(wb, RUN_INFO_SPECS[args.family], results_root=args.results_root)
    wb.save(args.xlsx)
    print(f"[run_info] {args.family} -> {args.xlsx}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
